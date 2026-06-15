from __future__ import annotations

from io import BytesIO

from docx import Document
from openpyxl import Workbook

from app.modules.ai.parsers import append_tables_markdown, extract_tables_from_bytes


def _build_xlsx_bytes(sheet_title: str, rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = sheet_title
    for row in rows:
        worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_docx_bytes(rows: list[list[str]]) -> bytes:
    document = Document()
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_pdf_bytes(table_lines: list[str]) -> bytes:
    text_ops = " T* ".join(f"({line}) Tj" for line in table_lines)
    content_stream = f"BT /F1 12 Tf 72 720 Td 14 TL {text_ops} ET".encode("ascii")
    objects = [
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
        (
            b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>\nendobj\n"
        ),
        b"4 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
        (
            b"5 0 obj\n<< /Length "
            + str(len(content_stream)).encode("ascii")
            + b" >>\nstream\n"
            + content_stream
            + b"\nendstream\nendobj\n"
        ),
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for obj in objects:
        offsets.append(len(output))
        output += obj
    xref_position = len(output)
    output += b"xref\n0 6\n0000000000 65535 f \n"
    for offset in offsets:
        output += f"{offset:010d} 00000 n \n".encode("ascii")
    output += (
        b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n"
        + str(xref_position).encode("ascii")
        + b"\n%%EOF\n"
    )
    return bytes(output)


def test_xlsx_extracts_headers_rows_columns_and_markdown() -> None:
    content = _build_xlsx_bytes("财务", [["合同编号", "金额"], ["KU-001", "100"]])

    tables = extract_tables_from_bytes(content, "xlsx")

    assert len(tables) == 1
    assert tables[0]["title"] == "财务"
    assert tables[0]["headers"] == ["合同编号", "金额"]
    assert tables[0]["rows"] == [["KU-001", "100"]]
    assert tables[0]["columns"] == [["KU-001"], ["100"]]
    assert "| 合同编号 | 金额 |" in str(tables[0]["markdown"])


def test_docx_extracts_table_structure() -> None:
    content = _build_docx_bytes([["姓名", "部门"], ["张三", "研发"]])

    tables = extract_tables_from_bytes(content, "docx")

    assert len(tables) == 1
    assert tables[0]["headers"] == ["姓名", "部门"]
    assert tables[0]["rows"] == [["张三", "研发"]]


def test_pdf_extracts_table_like_text_lines() -> None:
    content = _build_pdf_bytes(["Name | Amount", "Alice | 100"])

    tables = extract_tables_from_bytes(content, "pdf")

    assert len(tables) == 1
    assert tables[0]["headers"] == ["Name", "Amount"]
    assert tables[0]["rows"] == [["Alice", "100"]]


def test_append_tables_markdown_respects_character_limit() -> None:
    content = _build_xlsx_bytes("表", [["列A", "列B"], ["值A", "值B"]])
    tables = extract_tables_from_bytes(content, "xlsx")

    combined = append_tables_markdown("正文", tables, max_chars=16)

    assert combined.startswith("正文")
    assert len(combined) == 16
